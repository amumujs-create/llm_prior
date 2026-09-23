"""Read-only byte and workbook audit for the official E16-B CCPP source."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import struct
from collections import Counter
from pathlib import Path

import openpyxl


EXPECTED_COLUMNS = ("AT", "V", "AP", "RH", "PE")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def row_bytes(values: tuple[float, ...]) -> bytes:
    return b"".join(struct.pack(">d", float(value)) for value in values)


def digest_rows(rows: list[bytes], ordered: bool) -> str:
    digest = hashlib.sha256()
    iterator = rows if ordered else sorted(rows)
    for row in iterator:
        digest.update(row)
    return digest.hexdigest()


def sheet_audit(workbook: openpyxl.Workbook, title: str) -> dict:
    sheet = workbook[title]
    rows = list(sheet.iter_rows(values_only=True))
    header = tuple(rows[0])
    values = [tuple(float(value) for value in row) for row in rows[1:]]
    flattened = [value for row in values for value in row]
    encoded_rows = [row_bytes(row) for row in values]
    counts = Counter(encoded_rows)
    duplicate_groups = sum(count > 1 for count in counts.values())
    duplicate_excess_rows = sum(count - 1 for count in counts.values() if count > 1)
    return {
        "sheet_name": title,
        "shape": [len(values), len(header)],
        "columns": list(header),
        "missing_cells": sum(value is None for row in rows[1:] for value in row),
        "nonfinite_cells": sum(not math.isfinite(value) for value in flattened),
        "ordered_row_sha256": digest_rows(encoded_rows, ordered=True),
        "row_multiset_sha256": digest_rows(encoded_rows, ordered=False),
        "duplicate_row_groups": duplicate_groups,
        "duplicate_excess_rows": duplicate_excess_rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--zip", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--ods", type=Path, required=True)
    parser.add_argument("--readme", type=Path, required=True)
    parser.add_argument("--readme-backup", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    sheets = [sheet_audit(workbook, title) for title in workbook.sheetnames]
    canonical = sheets[0]
    structural_pass = (
        len(sheets) == 5
        and all(item["shape"] == [9568, 5] for item in sheets)
        and all(tuple(item["columns"]) == EXPECTED_COLUMNS for item in sheets)
        and all(item["missing_cells"] == 0 for item in sheets)
        and all(item["nonfinite_cells"] == 0 for item in sheets)
        and all(item["row_multiset_sha256"] == canonical["row_multiset_sha256"] for item in sheets)
    )

    artifact = {
        "protocol": "E16-B CCPP official source byte and workbook audit v1",
        "source_status": "PASS" if structural_pass else "REJECT",
        "source_identity": {
            "official_doi": "10.24432/C5002N",
            "official_archive_url": "https://archive.ics.uci.edu/static/public/294/combined%2Bcycle%2Bpower%2Bplant.zip",
            "license": "CC BY 4.0",
            "zip_sha256": sha256(args.zip),
            "zip_byte_size": args.zip.stat().st_size,
            "extracted_file_sha256": {
                "Folds5x2_pp.xlsx": sha256(args.xlsx),
                "Folds5x2_pp.ods": sha256(args.ods),
                "Readme.txt": sha256(args.readme),
                "Readme.txt~": sha256(args.readme_backup),
            },
        },
        "workbook_audit": {
            "sheet_count": len(sheets),
            "canonical_sheet": workbook.sheetnames[0],
            "sheets": sheets,
            "all_sheets_same_row_multiset": all(
                item["row_multiset_sha256"] == canonical["row_multiset_sha256"] for item in sheets
            ),
            "canonical_duplicate_policy": "recorded only; no row removed or deduplicated",
        },
        "audit_scope": "source structure only; no support cutoff, split, target-derived prior, model, policy, or outcome was computed",
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(artifact, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"source_status": artifact["source_status"], "canonical_sheet": workbook.sheetnames[0]}, indent=2))


if __name__ == "__main__":
    main()
